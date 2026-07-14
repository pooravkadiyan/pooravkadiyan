#!/usr/bin/env python3
"""Advanced Search Techniques — simple PowerPoint for 7th grade."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x3D, 0x4A)
TEAL = RGBColor(0x1F, 0x8A, 0x7A)
ORANGE = RGBColor(0xF0, 0x7F, 0x3C)
LIGHT = RGBColor(0xF3, 0xF8, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xD4, 0xEB, 0xE4)
MUTED = RGBColor(0x4E, 0x66, 0x68)
DARK = RGBColor(0x1A, 0x2E, 0x32)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
IMG = Path(__file__).resolve().parent / "images"
TOTAL = 10


def font(run, size=20, bold=False, color=DARK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def round_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def textbox(slide, left, top, width, height, text, size=20, bold=False,
            color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    font(run, size=size, bold=bold, color=color)
    return box


def bullets(slide, left, top, width, height, items, size=20, color=DARK, spacing=12):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = "•  " + item
        font(run, size=size, color=color)
    return box


def header(slide, title):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, LIGHT)
    rect(slide, 0, 0, Inches(0.2), SLIDE_H, TEAL)
    rect(slide, 0, 0, SLIDE_W, Inches(1.05), NAVY)
    textbox(slide, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.55),
            title, size=30, bold=True, color=WHITE)


def footer(slide, page):
    textbox(slide, Inches(0.5), Inches(7.05), Inches(9), Inches(0.3),
            "Advanced Search Tips  ·  For Grade 7", size=12, color=MUTED)
    textbox(slide, Inches(11.0), Inches(7.05), Inches(1.8), Inches(0.3),
            f"{page} / {TOTAL}", size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def picture(slide, name, left, top, width, height):
    path = IMG / name
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 — Title
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    picture(s, "01-cover-search.png", Inches(6.6), Inches(1.2), Inches(6.2), Inches(4.0))
    rect(s, 0, Inches(5.7), SLIDE_W, Inches(1.8), TEAL)
    textbox(s, Inches(0.55), Inches(1.8), Inches(6), Inches(0.5),
            "Grade 7 Skills", size=16, bold=True, color=ORANGE)
    textbox(s, Inches(0.55), Inches(2.4), Inches(6), Inches(1.4),
            "Advanced Search\nTechniques", size=40, bold=True, color=WHITE)
    textbox(s, Inches(0.55), Inches(4.2), Inches(5.8), Inches(0.8),
            "Find better answers online — the smart way!",
            size=18, color=SOFT)
    textbox(s, Inches(0.55), Inches(6.25), Inches(12), Inches(0.4),
            "Simple tips to search like a pro", size=16, color=WHITE)

    # 2 — Agenda (What we'll learn)
    s = prs.slides.add_slide(blank)
    header(s, "What We’ll Learn Today")
    tips = [
        ("1", "Why better searching helps"),
        ("2", "AND, OR, and NOT words"),
        ("3", "Quotes and star tricks"),
        ("4", "Cool Google tools"),
        ("5", "Where to search"),
        ("6", "Follow the book trail"),
        ("7", "Is this website true?"),
        ("8", "Try it yourself!"),
    ]
    for i, (num, label) in enumerate(tips):
        col = i % 2
        row = i // 2
        x = Inches(0.5 + col * 6.4)
        y = Inches(1.4 + row * 1.25)
        round_rect(s, x, y, Inches(6.0), Inches(1.05), WHITE)
        textbox(s, x + Inches(0.3), y + Inches(0.28), Inches(0.7), Inches(0.5),
                num, size=26, bold=True, color=ORANGE)
        textbox(s, x + Inches(1.1), y + Inches(0.32), Inches(4.6), Inches(0.5),
                label, size=20, color=DARK)
    footer(s, 2)

    # 3 — Why it matters
    s = prs.slides.add_slide(blank)
    header(s, "Why Search Skills Matter")
    picture(s, "02-why-search.png", Inches(7.0), Inches(1.35), Inches(5.7), Inches(5.2))
    round_rect(s, Inches(0.5), Inches(1.4), Inches(6.2), Inches(5.1), WHITE)
    textbox(s, Inches(0.8), Inches(1.7), Inches(5.6), Inches(0.5),
            "Searching isn’t just typing words!", size=20, bold=True, color=TEAL)
    bullets(s, Inches(0.8), Inches(2.4), Inches(5.6), Inches(3.8), [
        "Google shows millions of results.",
        "Many of them are ads or junk.",
        "Smart search = better homework help.",
        "You save time and find real facts.",
        "This skill helps you forever!",
    ], size=18, spacing=14)
    footer(s, 3)

    # 4 — Boolean AND OR NOT
    s = prs.slides.add_slide(blank)
    header(s, "Magic Words: AND · OR · NOT")
    picture(s, "03-boolean.png", Inches(0.4), Inches(1.25), Inches(5.8), Inches(3.5))
    cards = [
        ("AND", "Puts ideas together", 'dogs AND training', TEAL),
        ("OR", "Gives you more choices", 'kids OR children', NAVY),
        ("NOT", "Takes something away", 'jaguar NOT car', ORANGE),
    ]
    for i, (op, meaning, example, color) in enumerate(cards):
        x = Inches(0.5 + i * 4.2)
        y = Inches(4.85)
        round_rect(s, x, y, Inches(4.0), Inches(1.7), WHITE)
        rect(s, x, y, Inches(4.0), Inches(0.45), color)
        textbox(s, x + Inches(0.2), y + Inches(0.05), Inches(3.6), Inches(0.35),
                op, size=18, bold=True, color=WHITE)
        textbox(s, x + Inches(0.2), y + Inches(0.55), Inches(3.6), Inches(0.35),
                meaning, size=14, color=MUTED)
        textbox(s, x + Inches(0.2), y + Inches(0.95), Inches(3.6), Inches(0.5),
                example, size=16, bold=True, color=DARK)
    # Side tip over image area right
    round_rect(s, Inches(6.5), Inches(1.3), Inches(6.3), Inches(3.3), WHITE)
    textbox(s, Inches(6.8), Inches(1.55), Inches(5.8), Inches(0.45),
            "Think of it like this:", size=18, bold=True, color=TEAL)
    bullets(s, Inches(6.8), Inches(2.2), Inches(5.8), Inches(2.2), [
        "AND = both words must show up",
        "OR = either word is okay",
        "NOT = don’t show me that!",
    ], size=18, spacing=16)
    footer(s, 4)

    # 5 — Quotes and stars
    s = prs.slides.add_slide(blank)
    header(s, "Quotes \" \" and the Star *")
    picture(s, "04-phrases.png", Inches(7.0), Inches(1.3), Inches(5.8), Inches(5.3))
    round_rect(s, Inches(0.5), Inches(1.35), Inches(6.2), Inches(2.4), WHITE)
    textbox(s, Inches(0.8), Inches(1.55), Inches(5.6), Inches(0.4),
            'Use quotes for exact words', size=20, bold=True, color=TEAL)
    textbox(s, Inches(0.8), Inches(2.15), Inches(5.6), Inches(1.2),
            '"global warming"\n\nThis keeps the words stuck together —\nin that exact order.',
            size=17, color=DARK)

    round_rect(s, Inches(0.5), Inches(4.0), Inches(6.2), Inches(2.4), WHITE)
    textbox(s, Inches(0.8), Inches(4.2), Inches(5.6), Inches(0.4),
            "Use a star * for word endings", size=20, bold=True, color=ORANGE)
    textbox(s, Inches(0.8), Inches(4.8), Inches(5.6), Inches(1.3),
            "teach*\n\nFinds: teach, teacher, teaching,\nteachers… all at once!",
            size=17, color=DARK)
    footer(s, 5)

    # 6 — Google tricks
    s = prs.slides.add_slide(blank)
    header(s, "Cool Google Tricks")
    picture(s, "05-operators.png", Inches(0.35), Inches(1.25), Inches(5.6), Inches(5.3))
    tricks = [
        ("site:.edu", "Search school websites"),
        ("filetype:pdf", "Find PDF files only"),
        ("intitle:", "Word must be in the title"),
        ("-word", "Hide a word you don’t want"),
    ]
    for i, (op, meaning) in enumerate(tricks):
        y = Inches(1.35 + i * 1.25)
        round_rect(s, Inches(6.2), y, Inches(6.6), Inches(1.1), WHITE)
        textbox(s, Inches(6.5), y + Inches(0.18), Inches(6.0), Inches(0.35),
                op, size=20, bold=True, color=TEAL)
        textbox(s, Inches(6.5), y + Inches(0.55), Inches(6.0), Inches(0.35),
                meaning, size=16, color=DARK)
    footer(s, 6)

    # 7 — Where to search
    s = prs.slides.add_slide(blank)
    header(s, "Where Should You Search?")
    picture(s, "06-tools.png", Inches(7.0), Inches(1.3), Inches(5.8), Inches(5.3))
    places = [
        ("Google", "Good for everyday questions"),
        ("Google Scholar", "Better for school research papers"),
        ("Library / databases", "Books & articles your school trusts"),
        ("Kid-safe sites", ".edu, museums, NASA, Britannica"),
    ]
    for i, (name, desc) in enumerate(places):
        y = Inches(1.4 + i * 1.25)
        round_rect(s, Inches(0.5), y, Inches(6.2), Inches(1.1), WHITE)
        textbox(s, Inches(0.8), y + Inches(0.18), Inches(5.6), Inches(0.35),
                name, size=20, bold=True, color=NAVY)
        textbox(s, Inches(0.8), y + Inches(0.55), Inches(5.6), Inches(0.35),
                desc, size=16, color=MUTED)
    footer(s, 7)

    # 8 — Follow good sources
    s = prs.slides.add_slide(blank)
    header(s, "Follow the Trail of Good Sources")
    picture(s, "07-citations.png", Inches(7.0), Inches(1.3), Inches(5.8), Inches(5.3))
    round_rect(s, Inches(0.5), Inches(1.4), Inches(6.2), Inches(5.1), WHITE)
    textbox(s, Inches(0.8), Inches(1.7), Inches(5.6), Inches(0.5),
            "Found one great article?", size=20, bold=True, color=TEAL)
    bullets(s, Inches(0.8), Inches(2.4), Inches(5.6), Inches(3.8), [
        "Look at its Sources or References.",
        "Those books and links can help too!",
        "In Google Scholar, click “Cited by.”",
        "That shows newer papers that used it.",
        "It’s like a treasure map of ideas!",
    ], size=18, spacing=14)
    footer(s, 8)

    # 9 — Can I trust this?
    s = prs.slides.add_slide(blank)
    header(s, "Stop! Can You Trust This?")
    picture(s, "08-evaluate.png", Inches(7.0), Inches(1.3), Inches(5.8), Inches(5.3))
    checks = [
        ("Who made it?", "A teacher, scientist, or just a random blog?"),
        ("Is it true?", "Does it give facts and sources?"),
        ("Is it fresh?", "Old news may not help anymore."),
        ("Why was it made?", "To teach you — or to sell you something?"),
    ]
    for i, (q, a) in enumerate(checks):
        y = Inches(1.35 + i * 1.25)
        round_rect(s, Inches(0.5), y, Inches(6.2), Inches(1.1), WHITE)
        textbox(s, Inches(0.8), y + Inches(0.15), Inches(5.6), Inches(0.35),
                q, size=18, bold=True, color=ORANGE)
        textbox(s, Inches(0.8), y + Inches(0.55), Inches(5.6), Inches(0.35),
                a, size=15, color=DARK)
    footer(s, 9)

    # 10 — Practice + closing
    s = prs.slides.add_slide(blank)
    header(s, "Your Turn — Try These!")
    examples = [
        ('"solar system" planets filetype:pdf', "Find school-style PDFs"),
        ("volcanoes AND eruption site:.edu", "Only education sites"),
        ("frogs OR amphibians -poison", "More results, skip poison frogs"),
        ('intitle:"water cycle" kids', "Title must say water cycle"),
    ]
    for i, (query, tip) in enumerate(examples):
        y = Inches(1.3 + i * 1.05)
        round_rect(s, Inches(0.5), y, Inches(12.3), Inches(0.9), WHITE)
        textbox(s, Inches(0.8), y + Inches(0.12), Inches(11.7), Inches(0.35),
                query, size=17, bold=True, color=TEAL)
        textbox(s, Inches(0.8), y + Inches(0.48), Inches(11.7), Inches(0.3),
                tip, size=14, color=MUTED)

    # Closing mini bar
    round_rect(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.15), NAVY)
    textbox(s, Inches(0.8), Inches(5.85), Inches(11.7), Inches(0.7),
            "Remember: Smart searches = better answers. You’ve got this!",
            size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(s, 10)

    out = Path(__file__).resolve().parent / "Advanced_Search_Techniques.pptx"
    prs.save(str(out))
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
